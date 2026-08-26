## ALl document processing utilities

import logging
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation

from socratic.config import settings
from socratic.models import Resource, ResourceType

logger = logging.getLogger(__name__)

## This funciton extracts the text from PDF files.
## It takes the pdf path as parameter and returns the object with extracted text
def process_pdf(pdf_path: Path) -> Resource:
    logger.info(f"Processing PDF: {pdf_path}")

    try:
        reader = PdfReader(pdf_path)
        text_parts = []
        page_spans = []

        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text.strip():
                text_parts.append(text)
                page_spans.append(f"p.{i+1}")

        content = "\n\n".join(text_parts)

        resource = Resource(
            id=f"pdf_{pdf_path.stem}",
            type=ResourceType.PDF,
            url=str(pdf_path),
            title=pdf_path.stem,
            span=f"pp.{page_spans[0]}-{page_spans[-1]}" if page_spans else None,
            content=content,
            metadata={"num_pages": len(reader.pages), "file_path": str(pdf_path)},
        )

        logger.info(f"Extracted {len(reader.pages)} pages from {pdf_path}")
        return resource

    except Exception as e:
        logger.error(f"Error processing PDF {pdf_path}: {e}")
        raise

## This function extracts the text from powerpoint slides (I dont remember whether or not they existed but it doesnt hurt to check)
def process_slides(pptx_path: Path) -> Resource:
    logger.info(f"Processing slides: {pptx_path}")

    try:
        prs = Presentation(pptx_path)
        text_parts = []
        slide_spans = []

        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)

            if slide_text:
                text_parts.append(f"[Slide {i+1}]\n" + "\n".join(slide_text))
                slide_spans.append(str(i + 1))

        content = "\n\n".join(text_parts)

        resource = Resource(
            id=f"slides_{pptx_path.stem}",
            type=ResourceType.SLIDE,
            url=str(pptx_path),
            title=pptx_path.stem,
            span=f"slides {slide_spans[0]}-{slide_spans[-1]}" if slide_spans else None,
            content=content,
            metadata={"num_slides": len(prs.slides), "file_path": str(pptx_path)},
        )

        logger.info(f"Extracted {len(prs.slides)} slides from {pptx_path}")
        return resource

    except Exception as e:
        logger.error(f"Error processing slides {pptx_path}: {e}")
        raise

## This function splits the text into overlapping chunks.
## It takes in the text to chunk, the size of each chunk in characters, and the overlap between chunks in characters as parameters
## and retuns a list of text chunks
def chunk_text(
    text: str, chunk_size: int | None = None, overlap: int | None = None
) -> list[str]:

    chunk_size = chunk_size or settings.chunk_size
    overlap = overlap or settings.chunk_overlap

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]

        ## Try to break at sentence boundary
        if end < len(text):
            last_period = chunk.rfind(".")
            last_newline = chunk.rfind("\n")
            break_point = max(last_period, last_newline)

            if break_point > chunk_size * 0.5: ## Only break if we're past halfway
                chunk = chunk[: break_point + 1]
                end = start + break_point + 1

        chunks.append(chunk.strip())
        start = end - overlap

    return chunks
